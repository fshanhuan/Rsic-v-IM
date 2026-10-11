# -*- coding: utf-8 -*-
"""
md2docx.py — 把 `docs/产品说明_v9_客户版.md` 转成排版好的 Word 文档。

支持本工程用到的 Markdown 子集：
  标题 # ## ###、段落、无序/有序列表、表格、围栏代码块、引用块、分隔线、
  行内 **加粗** 与 `代码`、图片 ![alt](path)。
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "docs", "产品说明_v9_客户版.md")
OUT = os.path.join(ROOT, "v9_处理器产品说明_客户版.docx")

FONT_CJK = "微软雅黑"
FONT_BODY_CJK = "宋体"
FONT_LATIN = "Times New Roman"
FONT_MONO = "Consolas"

ACCENT = RGBColor(0x1F, 0x3B, 0x63)
GREY = RGBColor(0x5A, 0x62, 0x70)


# ----------------------------------------------------------------- 低层工具
def set_run_font(run, cjk=FONT_BODY_CJK, latin=FONT_LATIN, size=None,
                 bold=None, italic=None, color=None, mono=False):
    if mono:
        cjk, latin = FONT_MONO, FONT_MONO
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph_or_cell, hex_fill):
    el = paragraph_or_cell._element
    pr = el.get_or_add_pPr() if el.tag.endswith("}p") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def cell_shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


# ----------------------------------------------------------------- 行内解析
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, base_cjk=FONT_BODY_CJK,
               base_color=None, base_bold=False, depth=0):
    """递归解析行内标记：**加粗**（其中还可再嵌 `代码`）与 `代码`。"""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(paragraph, tok[2:-2], size=size, base_cjk=FONT_CJK,
                       base_color=base_color or ACCENT, base_bold=True,
                       depth=depth + 1)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            set_run_font(r, size=size - 0.5, mono=True,
                         bold=base_bold,
                         color=RGBColor(0x8B, 0x2E, 0x2E))
        else:
            r = paragraph.add_run(tok)
            set_run_font(r, cjk=base_cjk, size=size, bold=base_bold,
                         color=base_color)


# ----------------------------------------------------------------- 解析 MD
def parse(md_text):
    lines = md_text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s:
            i += 1
            continue
        if s.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", buf))
            continue
        if re.fullmatch(r"-{3,}", s):
            blocks.append(("hr", None))
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            blocks.append(("table", rows))
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", buf))
            continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            blocks.append(("img", (m.group(1), m.group(2))))
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            buf = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                buf.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ul", buf))
            continue
        if re.match(r"^\d+\.\s+", s):
            buf = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                buf.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ol", buf))
            continue
        buf = [s]
        i += 1
        while i < len(lines):
            nx = lines[i].strip()
            if (not nx or nx.startswith(("#", "|", ">", "```"))
                    or re.match(r"^[-*]\s+", nx) or re.match(r"^\d+\.\s+", nx)
                    or re.fullmatch(r"-{3,}", nx)
                    or re.match(r"^!\[.*?\]\(.*?\)$", nx)):
                break
            buf.append(nx)
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


def split_table(rows):
    cells = []
    for r in rows:
        if re.fullmatch(r"\|[\s:|-]+\|", r):
            continue
        parts = [c.strip() for c in r.strip().strip("|").split("|")]
        cells.append(parts)
    return cells


# ----------------------------------------------------------------- 主流程
def build():
    with open(MD, encoding="utf-8") as fh:
        blocks = parse(fh.read())

    doc = Document()

    # 页面设置
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    # 正文默认字体
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = FONT_LATIN
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY_CJK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    # 页脚页码
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("BigBird v9 产品说明   ·   第 ")
    set_run_font(r, size=8.5, color=GREY)
    add_page_field(fp)
    r2 = fp.add_run(" 页")
    set_run_font(r2, size=8.5, color=GREY)

    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(10)
            add_inline(p, payload, size=19, base_cjk=FONT_CJK,
                       base_color=ACCENT, base_bold=True)
        elif kind in ("h2", "h3", "h4"):
            level = {"h2": 1, "h3": 2, "h4": 3}[kind]
            p = doc.add_paragraph(style="Heading %d" % level)
            size = {1: 14.5, 2: 12.5, 3: 11.5}[level]
            p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, payload, size=size, base_cjk=FONT_CJK,
                       base_color=ACCENT if level == 1 else RGBColor(0x24, 0x4A, 0x77),
                       base_bold=True)
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline(p, payload)
        elif kind == "ul":
            for item in payload:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, item)
        elif kind == "ol":
            for idx, item in enumerate(payload, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.85)
                p.paragraph_format.first_line_indent = Cm(-0.85)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run("%d. " % idx)
                set_run_font(r, size=10.5, bold=False)
                add_inline(p, item)
        elif kind == "quote":
            for item in payload:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.4)
                p.paragraph_format.space_after = Pt(2)
                shade(p, "F2F5FA")
                add_inline(p, item, size=10, base_color=RGBColor(0x33, 0x3D, 0x4D))
        elif kind == "code":
            for item in payload:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.left_indent = Cm(0.35)
                shade(p, "F5F5F5")
                r = p.add_run(item if item else " ")
                set_run_font(r, size=9, mono=True,
                             color=RGBColor(0x1A, 0x1A, 0x1A))
            tail = doc.add_paragraph()
            tail.paragraph_format.space_after = Pt(6)
        elif kind == "img":
            alt, path = payload
            full = os.path.normpath(os.path.join(os.path.dirname(MD), path))
            if os.path.exists(full):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(4)
                p.add_run().add_picture(full, width=Cm(16.4))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(10)
                cr = cap.add_run(alt)
                set_run_font(cr, size=9, color=GREY, italic=True)
            else:
                print("WARN missing image:", full)
        elif kind == "table":
            rows = split_table(payload)
            if not rows:
                continue
            ncol = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for ri, cells in enumerate(rows):
                row = table.add_row()
                if ri == 0:
                    repeat_header(row)
                for ci in range(ncol):
                    txt = cells[ci] if ci < len(cells) else ""
                    cell = row.cells[ci]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.line_spacing = 1.05
                    if ri == 0:
                        cell_shade(cell, "E8EFF8")
                        add_inline(p, txt, size=9.5, base_cjk=FONT_CJK,
                                   base_color=ACCENT, base_bold=True)
                    else:
                        add_inline(p, txt, size=9)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(8)
        elif kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            ppr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C9D2E0")
            pbdr.append(bottom)
            ppr.append(pbdr)

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
