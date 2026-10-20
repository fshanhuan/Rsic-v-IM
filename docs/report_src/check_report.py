# -*- coding: utf-8 -*-
"""Verify the generated product report: structure, key numbers, stray markup."""
import os
import re
import sys
from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
name = sys.argv[1] if len(sys.argv) > 1 else "产品报告_v17.docx"
path = os.path.join(ROOT, "docs", name)
print("checking", name)
doc = Document(path)

texts = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            texts.extend(p.text for p in c.paragraphs)
alltext = "\n".join(texts)

out = []
out.append("paragraphs=%d tables=%d images=%d chars=%d"
           % (len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes), len(alltext)))

heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
out.append("\nHeadings (%d):" % len(heads))
for h in heads:
    out.append("   " + h)

bad = [s for s in ("**", "TODO", "XXX", "lorem", "None") if s in alltext]
out.append("\nstray tokens: %s" % (bad or "none"))

facts = ["128.82", "2.576", "1400", "38/38", "e714", "0.048", "36 150", "34 058",
         "66.5", "2.335", "2.408", "22 681", "468 266", "471 384", "0xf7b9ad39",
         "0x00011197", "10.868", "543407157", "50 MHz", "115200", "1.245", "1.323",
         "36 990", "0.645", "1.187", "16 384", "32 768", "4 731", "1 543",
         "xc7z020", "BTB/PHT 256", "GHR 8", "fence.i", "IROM_APP_MIN"]
missing = [f for f in facts if f not in alltext]
out.append("missing facts: %s" % (missing or "none"))

caps = sorted({int(n) for n in re.findall(r"表\s?(\d+)　", alltext)})
out.append("table captions: %d..%d count=%d" % (caps[0], caps[-1], len(caps)))
out.append("figure captions: %s" % re.findall(r"图\s?(\d+)　", alltext))

out.append("\ntables: index, cols x rows, header cell")
for i, t in enumerate(doc.tables):
    hdr = t.rows[0].cells[0].text.strip().replace("\n", " ")[:16]
    out.append("  %2d  %dx%d  %s" % (i, len(t.columns), len(t.rows), hdr))

report = "\n".join(out)
with open(os.path.join(ROOT, "docs", "report_src", "check_report.txt"), "w",
          encoding="utf-8") as f:
    f.write(report)
print("wrote check_report.txt; missing=%d stray=%d" % (len(missing), len(bad)))
if missing or bad:
    sys.exit(1)
